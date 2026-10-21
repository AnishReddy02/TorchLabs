from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).parent / "vendor"))

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).parent
SOURCE = ROOT / "Torch_Lab_Final_working.pptx"
OUTPUT = ROOT / "Torch_Labs_Final_With_Product_Workflow.pptx"
ASSETS = ROOT / "assets"

NAVY = RGBColor(0x17, 0x2B, 0x43)
BLUE = RGBColor(0x35, 0x6C, 0xD5)
GREEN = RGBColor(0x3B, 0xD1, 0x8B)
MID = RGBColor(0x5D, 0x6B, 0x7A)
PALE = RGBColor(0xE4, 0xE9, 0xEF)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def crop(source: str, output: str, box: tuple[int, int, int, int]) -> Path:
    target = ASSETS / output
    with Image.open(ASSETS / source) as image:
        image.crop(box).save(target, quality=95)
    return target


def text(slide, value, x, y, w, h, size, color=NAVY, font="Segoe UI", bold=False,
         align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def picture_frame(slide, path: Path, x, y, w, h):
    frame = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x - 0.04), Inches(y - 0.04),
        Inches(w + 0.08), Inches(h + 0.08)
    )
    frame.fill.solid()
    frame.fill.fore_color.rgb = NAVY
    frame.line.color.rgb = PALE
    frame.line.width = Pt(1)
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))


def add_header(slide, label, title, number, subtitle):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE
    text(slide, label.upper(), 0.8, 0.34, 8.8, 0.3, 11.25, BLUE)
    text(slide, number, 12.0, 0.34, 0.65, 0.3, 11.25, BLUE, align=PP_ALIGN.RIGHT)
    text(slide, title, 0.8, 0.82, 11.8, 0.62, 30, NAVY, font="Georgia")
    text(slide, subtitle, 0.8, 1.46, 11.7, 0.42, 15, MID)


def add_full_product_slide(prs, label, title, number, subtitle, image_path):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    add_header(slide, label, title, number, subtitle)
    picture_frame(slide, image_path, 0.8, 2.02, 11.75, 4.95)
    return slide


prs = Presentation(SOURCE)

# Crop to the meaningful product area so screenshots remain legible on a slide.
asset_crop = crop("torch-asset-ready.png", "deck-asset-ready.png", (145, 185, 1295, 590))
mission_crop = crop("torch-mission-control.png", "deck-mission-control.png", (150, 0, 1295, 655))
approval_crop = crop("torch-approval.png", "deck-approval.png", (145, 70, 1295, 665))
run_crop = crop("torch-run-failure.png", "deck-run-failure.png", (145, 75, 1295, 965))
report_crop = crop("torch-agent-report.png", "deck-agent-report.png", (145, 70, 1295, 610))

# Replace the explicit placeholder on the existing solution slide.
solution = prs.slides[3]
for shape in list(solution.shapes):
    if "[Add screenshot" in getattr(shape, "text", ""):
        shape._element.getparent().remove(shape._element)
picture_frame(solution, asset_crop, 2.63, 3.05, 8.08, 2.85)

new_slides = []
new_slides.append(add_full_product_slide(
    prs, "PRODUCT · 1 OF 4", "The operator defines the mission and the limits", "05",
    "Choose the asset, deadline, failure policy, and objective before the agent acts.", mission_crop
))
new_slides.append(add_full_product_slide(
    prs, "PRODUCT · 2 OF 4", "Software changes stop for human approval", "06",
    "Torch shows the exact asset, workflow, bundle, reason, and SHA before deployment.", approval_crop
))

execution = prs.slides.add_slide(prs.slide_layouts[0])
add_header(execution, "PRODUCT · 3 OF 4", "Existing workflows run on real hardware", "07",
           "The node streams stdout, persists every step, attempts cleanup, and collects artifacts.")
text(execution, "PREPARE", 0.8, 2.28, 2.7, 0.28, 10, BLUE, bold=True)
text(execution, "Inspect build", 0.8, 2.57, 2.7, 0.36, 19, NAVY, font="Georgia")
text(execution, "TEST", 0.8, 3.25, 2.7, 0.28, 10, BLUE, bold=True)
text(execution, "Localization fails", 0.8, 3.54, 2.9, 0.36, 19, NAVY, font="Georgia")
text(execution, "CLEANUP", 0.8, 4.22, 2.7, 0.28, 10, BLUE, bold=True)
text(execution, "Still runs", 0.8, 4.51, 2.7, 0.36, 19, NAVY, font="Georgia")
text(execution, "✓", 3.33, 2.57, 0.35, 0.36, 18, GREEN, bold=True)
text(execution, "×", 3.33, 3.54, 0.35, 0.36, 20, RGBColor(0xE8, 0x5D, 0x68), bold=True)
text(execution, "✓", 3.33, 4.51, 0.35, 0.36, 18, GREEN, bold=True)
picture_frame(execution, run_crop, 4.0, 2.05, 8.55, 4.95)
new_slides.append(execution)

new_slides.append(add_full_product_slide(
    prs, "PRODUCT · 4 OF 4", "The agent explains the failure from actual evidence", "08",
    "The model sees bounded Torch logs, step results, machine context, and artifact metadata—not SSH.", report_crop
))

# Move the newly created slides directly after the original solution slide.
slide_ids = prs.slides._sldIdLst
new_slide_ids = list(slide_ids)[-len(new_slides):]
for element in new_slide_ids:
    slide_ids.remove(element)
for offset, element in enumerate(new_slide_ids):
    slide_ids.insert(4 + offset, element)

# Existing slides 5–8 now follow the four-slide product sequence.
for slide, number in zip(list(prs.slides)[8:], ["09", "10", "11", "12"]):
    for shape in slide.shapes:
        if getattr(shape, "text", "").strip() in {"05", "06", "07", "08"}:
            shape.text_frame.paragraphs[0].runs[0].text = number

prs.save(OUTPUT)
print(OUTPUT)
